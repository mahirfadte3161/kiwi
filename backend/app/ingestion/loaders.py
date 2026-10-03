from pathlib import Path
from .cleaner import clean_text


def load_pdf(path: Path) -> tuple[list[tuple[int, str]], int]:
    import fitz
    pages, ocr_count = [], 0
    with fitz.open(path) as document:
        for number, page in enumerate(document, 1):
            text = clean_text(page.get_text("text"))
            if len(text) < 20:
                ocr_count += 1
                try:
                    from PIL import Image
                    import pytesseract
                    import io
                    image = Image.open(io.BytesIO(page.get_pixmap(matrix=fitz.Matrix(2, 2)).tobytes("png")))
                    text = clean_text(pytesseract.image_to_string(image))
                except Exception:
                    text = text
            if text:
                pages.append((number, text))
    return pages, ocr_count


def load_pptx(path: Path) -> tuple[list[tuple[int, str]], int]:
    from pptx import Presentation
    presentation = Presentation(path)
    pages = []
    for number, slide in enumerate(presentation.slides, 1):
        parts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        notes = getattr(slide, "notes_slide", None)
        if notes and notes.notes_text_frame:
            parts.extend(shape.text for shape in notes.notes_text_frame.paragraphs if shape.text.strip())
        text = clean_text("\n".join(parts))
        if text:
            pages.append((number, text))
    return pages, 0
